"""Export service for generating PDF and Excel documents."""

import html
import io
import os
import re
import unicodedata
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    HRFlowable,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


def sanitize_filename(name: str | None) -> str:
    """Sanitize user-supplied filename for safe use in HTTP headers and filesystems.

    Strips or replaces quotes, backslashes, control characters, path separators,
    and semicolons. Ensures the resulting string is non-empty and safe.
    """
    if not name:
        return "document"

    # Strip any directory path components
    base_name = os.path.basename(str(name).strip())

    # Normalize unicode to ASCII representation where possible
    normalized = (
        unicodedata.normalize("NFKD", base_name)
        .encode("ascii", "ignore")
        .decode("ascii")
    )
    if not normalized:
        normalized = base_name

    # Replace quotes, backslashes, slashes, control chars, semicolons, and other unsafe chars
    sanitized = re.sub(r'[\r\n\t\x00-\x1f\x7f\\/"\';:*?<>|`]', "_", normalized)

    # Keep only alphanumeric, dot, hyphen, underscore, and space
    sanitized = re.sub(r"[^\w.\- ]", "_", sanitized)

    # Collapse multiple consecutive underscores
    sanitized = re.sub(r"_+", "_", sanitized)

    # Strip leading/trailing dots, spaces, and underscores
    sanitized = sanitized.strip(". _")

    if not sanitized:
        return "document"

    return sanitized


def format_file_size(size_bytes: int | None) -> str:
    """Format file size in bytes to a human-readable string."""
    if size_bytes is None:
        return "N/A"
    if size_bytes < 1024:
        return f"{size_bytes} B"
    elif size_bytes < 1024 * 1024:
        return f"{size_bytes / 1024:.1f} KB"
    elif size_bytes < 1024 * 1024 * 1024:
        return f"{size_bytes / (1024 * 1024):.1f} MB"
    else:
        return f"{size_bytes / (1024 * 1024 * 1024):.1f} GB"


def format_confidence(confidence: float | None) -> str:
    """Format confidence score as a percentage string."""
    if confidence is None:
        return "N/A"
    # If stored as 0.0 - 1.0
    val = confidence * 100.0 if confidence <= 1.0 else confidence
    return f"{val:.1f}%"


class _NumberedCanvas(canvas.Canvas):
    """Canvas that performs a two-pass render to display total page count."""

    def __init__(self, *args: Any, **kwargs: Any) -> None:
        super().__init__(*args, **kwargs)
        self._saved_page_states: list[dict[str, Any]] = []

    def showPage(self) -> None:
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self) -> None:
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count: int) -> None:
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        # Top page boundary: letter is (612, 792)
        # Footer text and divider rule
        self.setStrokeColor(colors.HexColor("#E2E8F0"))
        self.setLineWidth(0.5)
        self.line(40, 36, 572, 36)

        left_text = "DocuMind AI — Document Analysis & Export"
        right_text = f"Page {self._pageNumber} of {page_count}"
        self.drawString(40, 24, left_text)
        self.drawRightString(572, 24, right_text)
        self.restoreState()


def generate_pdf(document: Any) -> bytes:
    """Generate a ReportLab PDF export for the given document.

    Includes:
      - DocuMind AI header
      - Original filename and file type
      - Upload date
      - Metadata section (page count, word count, human-readable file size)
      - AI Analysis section (document_type, summary, classification_confidence)
        if processing_status is "completed", otherwise "Not yet processed"
      - Extracted text (up to ~3000 chars with truncation notice if applicable,
        or 'No text content available' if empty)

    Returns PDF bytes.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=48,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocuMindTitle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#0F172A"),
        spaceAfter=2,
    )
    subtitle_style = ParagraphStyle(
        "DocuMindSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748B"),
        spaceAfter=8,
    )
    section_heading_style = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#1E293B"),
        spaceBefore=10,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "BodyDark",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#334155"),
    )
    cell_label_style = ParagraphStyle(
        "CellLabel",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#475569"),
    )
    cell_value_style = ParagraphStyle(
        "CellValue",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#0F172A"),
    )
    notice_style = ParagraphStyle(
        "NoticeStyle",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#64748B"),
    )
    text_content_style = ParagraphStyle(
        "TextContent",
        parent=styles["Normal"],
        fontName="Courier",
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#1E293B"),
    )

    story: list[Any] = []

    # 1. Header
    story.append(Paragraph("DocuMind AI", title_style))
    story.append(
        Paragraph("Document Analysis & Intelligence Report", subtitle_style)
    )
    story.append(
        HRFlowable(
            width="100%",
            thickness=1.5,
            color=colors.HexColor("#3B82F6"),
            spaceBefore=0,
            spaceAfter=12,
        )
    )

    # 2. Document Overview & Metadata Table
    story.append(Paragraph("Document Overview", section_heading_style))

    upload_date_str = (
        document.created_at.strftime("%Y-%m-%d %H:%M:%S UTC")
        if getattr(document, "created_at", None)
        else "N/A"
    )

    meta_table_data = [
        [
            Paragraph("Original Filename", cell_label_style),
            Paragraph(
                html.escape(document.original_filename or "N/A"), cell_value_style
            ),
        ],
        [
            Paragraph("File Type", cell_label_style),
            Paragraph(
                html.escape((document.file_type or "N/A").upper()),
                cell_value_style,
            ),
        ],
        [
            Paragraph("Upload Date", cell_label_style),
            Paragraph(html.escape(upload_date_str), cell_value_style),
        ],
        [
            Paragraph("File Size", cell_label_style),
            Paragraph(
                html.escape(format_file_size(document.file_size_bytes)),
                cell_value_style,
            ),
        ],
        [
            Paragraph("Page Count", cell_label_style),
            Paragraph(
                str(document.page_count)
                if document.page_count is not None
                else "N/A",
                cell_value_style,
            ),
        ],
        [
            Paragraph("Word Count", cell_label_style),
            Paragraph(
                f"{document.word_count:,}"
                if document.word_count is not None
                else "N/A",
                cell_value_style,
            ),
        ],
        [
            Paragraph("Processing Status", cell_label_style),
            Paragraph(
                html.escape(document.processing_status or "pending"),
                cell_value_style,
            ),
        ],
    ]

    # Printable width: 612 - 80 = 532
    meta_table = Table(meta_table_data, colWidths=[140, 392])
    meta_table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
            ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#1E293B")),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ])
    )
    story.append(meta_table)
    story.append(Spacer(1, 10))

    # 3. AI Analysis Section
    story.append(Paragraph("AI Analysis", section_heading_style))

    is_completed = (
        getattr(document, "processing_status", "").lower() == "completed"
    )

    if is_completed:
        conf_str = format_confidence(
            getattr(document, "classification_confidence", None)
        )
        ai_table_data = [
            [
                Paragraph("Document Type", cell_label_style),
                Paragraph(
                    html.escape(document.document_type or "OTHER"),
                    cell_value_style,
                ),
            ],
            [
                Paragraph("Classification Confidence", cell_label_style),
                Paragraph(html.escape(conf_str), cell_value_style),
            ],
            [
                Paragraph("Summary", cell_label_style),
                Paragraph(
                    html.escape(document.summary or "No summary available.")
                    .replace("\n", "<br/>"),
                    cell_value_style,
                ),
            ],
        ]
        ai_table = Table(ai_table_data, colWidths=[140, 392])
        ai_table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F0FDF4")),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#BBF7D0")),
            ])
        )
        story.append(ai_table)
    else:
        pending_table_data = [
            [
                Paragraph("Status", cell_label_style),
                Paragraph("Not yet processed", cell_value_style),
            ],
            [
                Paragraph("Note", cell_label_style),
                Paragraph(
                    "This document has not been processed by the AI pipeline yet. "
                    "Run AI processing to generate classification, confidence score, and summary.",
                    notice_style,
                ),
            ],
        ]
        pending_table = Table(pending_table_data, colWidths=[140, 392])
        pending_table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FFFBEB")),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#FDE68A")),
            ])
        )
        story.append(pending_table)

    story.append(Spacer(1, 10))

    # 4. Extracted Text Section
    story.append(Paragraph("Extracted Text", section_heading_style))

    raw_text = getattr(document, "raw_text", None) or ""
    stripped_text = raw_text.strip()

    if not stripped_text:
        story.append(
            Paragraph("<i>No text content available</i>", notice_style)
        )
    else:
        TEXT_LIMIT = 3000
        truncated = len(stripped_text) > TEXT_LIMIT
        display_text = stripped_text[:TEXT_LIMIT]

        escaped_display = html.escape(display_text).replace("\n", "<br/>")
        if truncated:
            truncation_note = (
                f"<br/><br/><b>[Content truncated — showing first {TEXT_LIMIT:,} of "
                f"{len(stripped_text):,} characters. For complete text, please export to Excel.]</b>"
            )
            escaped_display += truncation_note

        text_block = [
            [Paragraph(escaped_display, text_content_style)],
        ]
        text_table = Table(text_block, colWidths=[532])
        text_table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("LEFTPADDING", (0, 0), (-1, -1), 10),
                ("RIGHTPADDING", (0, 0), (-1, -1), 10),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ])
        )
        story.append(text_table)

    # Build the document
    doc.build(story, canvasmaker=_NumberedCanvas)
    buffer.seek(0)
    return buffer.getvalue()


def generate_excel(document: Any) -> bytes:
    """Generate an Excel (.xlsx) workbook export for the given document.

    Contains three sheets:
      1. 'Overview': Original filename, system filename, file type, upload date,
         size, page count, word count, processing status.
      2. 'AI Analysis': document_type, summary, classification_confidence (as
         a percentage). If processing_status != "completed", shows
         "Document not yet processed" across this sheet.
      3. 'Full Text': The complete raw_text. Slices text into sequential chunks
         of 30,000 characters each down rows in Column A to respect Excel's
         32,767 character-per-cell limit.

    Returns .xlsx bytes.
    """
    wb = Workbook()

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="1E293B", end_color="1E293B", fill_type="solid"
    )
    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    label_font = Font(name="Calibri", size=10, bold=True, color="334155")
    value_font = Font(name="Calibri", size=10, color="0F172A")
    notice_font = Font(
        name="Calibri", size=11, bold=True, italic=True, color="B45309"
    )

    thin_border_side = Side(border_style="thin", color="E2E8F0")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )
    alt_fill = PatternFill(
        start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
    )

    # -------------------------------------------------------------
    # Sheet 1: Overview
    # -------------------------------------------------------------
    ws_overview = wb.active
    ws_overview.title = "Overview"
    ws_overview.views.sheetView[0].showGridLines = True

    # Title
    ws_overview.cell(row=1, column=1, value="DocuMind AI — Document Overview")
    ws_overview.cell(row=1, column=1).font = title_font

    # Headers
    ws_overview.cell(row=3, column=1, value="Attribute")
    ws_overview.cell(row=3, column=2, value="Value")
    for col in (1, 2):
        c = ws_overview.cell(row=3, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = cell_border

    upload_date_str = (
        document.created_at.strftime("%Y-%m-%d %H:%M:%S UTC")
        if getattr(document, "created_at", None)
        else "N/A"
    )

    overview_rows = [
        ("Original Filename", document.original_filename or "N/A"),
        ("System Filename", document.filename or "N/A"),
        ("File Type", (document.file_type or "N/A").upper()),
        ("Upload Date", upload_date_str),
        ("File Size", format_file_size(document.file_size_bytes)),
        (
            "Page Count",
            document.page_count if document.page_count is not None else "N/A",
        ),
        (
            "Word Count",
            document.word_count if document.word_count is not None else "N/A",
        ),
        ("Processing Status", document.processing_status or "pending"),
    ]

    for idx, (attr, val) in enumerate(overview_rows, start=4):
        c_attr = ws_overview.cell(row=idx, column=1, value=attr)
        c_val = ws_overview.cell(row=idx, column=2, value=val)

        c_attr.font = label_font
        c_val.font = value_font
        c_attr.border = cell_border
        c_val.border = cell_border

        if idx % 2 == 0:
            c_attr.fill = alt_fill
            c_val.fill = alt_fill

    ws_overview.column_dimensions["A"].width = 24
    ws_overview.column_dimensions["B"].width = 50

    # -------------------------------------------------------------
    # Sheet 2: AI Analysis
    # -------------------------------------------------------------
    ws_ai = wb.create_sheet(title="AI Analysis")
    ws_ai.views.sheetView[0].showGridLines = True

    ws_ai.cell(row=1, column=1, value="DocuMind AI — AI Analysis")
    ws_ai.cell(row=1, column=1).font = title_font

    is_completed = (
        getattr(document, "processing_status", "").lower() == "completed"
    )

    if not is_completed:
        # Prompt: show "Document not yet processed" across this sheet if processing_status isn't "completed"
        ws_ai.cell(row=3, column=1, value="Status")
        ws_ai.cell(row=3, column=2, value="Document not yet processed")
        for col in (1, 2):
            c = ws_ai.cell(row=3, column=col)
            c.font = header_font
            c.fill = PatternFill(
                start_color="92400E", end_color="92400E", fill_type="solid"
            )
            c.border = cell_border

        c_msg = ws_ai.cell(
            row=5,
            column=1,
            value="Document not yet processed. Run AI processing on this document to generate classification, confidence score, and summary.",
        )
        c_msg.font = notice_font

        ws_ai.column_dimensions["A"].width = 25
        ws_ai.column_dimensions["B"].width = 50
    else:
        ws_ai.cell(row=3, column=1, value="Metric / Field")
        ws_ai.cell(row=3, column=2, value="Result")
        for col in (1, 2):
            c = ws_ai.cell(row=3, column=col)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="left", vertical="center")
            c.border = cell_border

        conf_str = format_confidence(
            getattr(document, "classification_confidence", None)
        )

        ai_rows = [
            ("Processing Status", "completed"),
            ("Document Type", document.document_type or "OTHER"),
            ("Classification Confidence", conf_str),
            ("Summary", document.summary or "No summary available."),
        ]

        for idx, (field, val) in enumerate(ai_rows, start=4):
            c_field = ws_ai.cell(row=idx, column=1, value=field)
            c_val = ws_ai.cell(row=idx, column=2, value=val)

            c_field.font = label_font
            c_val.font = value_font
            c_field.border = cell_border
            c_val.border = cell_border
            c_field.alignment = Alignment(vertical="top")
            c_val.alignment = Alignment(vertical="top", wrap_text=True)

            if idx % 2 == 0:
                c_field.fill = alt_fill
                c_val.fill = alt_fill

        ws_ai.column_dimensions["A"].width = 28
        ws_ai.column_dimensions["B"].width = 80

    # -------------------------------------------------------------
    # Sheet 3: Full Text
    # -------------------------------------------------------------
    ws_text = wb.create_sheet(title="Full Text")
    ws_text.views.sheetView[0].showGridLines = True

    raw_text = getattr(document, "raw_text", None) or ""
    stripped_text = raw_text.strip()

    if not stripped_text:
        c = ws_text.cell(row=1, column=1, value="No text content available")
        c.font = notice_font
    else:
        # Excel cell limit is 32,767 chars. Chunk at 30,000 chars per row down column A.
        CHUNK_SIZE = 30000
        chunks = [
            raw_text[i : i + CHUNK_SIZE]
            for i in range(0, len(raw_text), CHUNK_SIZE)
        ]
        for row_idx, chunk in enumerate(chunks, start=1):
            cell = ws_text.cell(row=row_idx, column=1, value=chunk)
            cell.font = Font(name="Calibri", size=10, color="0F172A")
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws_text.column_dimensions["A"].width = 110

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
